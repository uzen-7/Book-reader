from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument

from core.document import Document, Page
from extraction.base_extractor import BaseExtractor


class DOCXExtractor(BaseExtractor):
    """Extract paragraph-based text from DOCX files."""

    def extract(self, file_path: str | Path) -> list[dict]:
        path = Path(file_path)
        doc = DocxDocument(str(path))
        pages: list[dict] = []
        paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text and paragraph.text.strip()]

        if not paragraphs:
            return [{"page_number": 1, "text": "", "source": "embedded", "image_path": None, "metadata": {}}]

        pages.append({
            "page_number": 1,
            "text": "\n\n".join(paragraphs),
            "source": "embedded",
            "image_path": None,
            "metadata": {},
        })
        return pages


def extract_docx_document(file_path: str | Path) -> Document:
    """Create a Document object from a DOCX file."""
    path = Path(file_path)
    pages = DOCXExtractor().extract(path)
    document_pages = [Page(page_number=item["page_number"], text=item["text"], source=item["source"], metadata=item["metadata"]) for item in pages]
    return Document(file_path=path, title=path.stem, pages=document_pages, document_type="docx")
