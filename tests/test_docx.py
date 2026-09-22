from pathlib import Path

from docx import Document

from extraction.docx_extractor import DOCXExtractor


def test_docx_extractor_extracts_paragraphs(tmp_path: Path):
    doc_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("This is a paragraph.")
    document.add_paragraph("This is another paragraph.")
    document.save(doc_path)

    result = DOCXExtractor().extract(doc_path)
    assert len(result) == 1
    assert "This is a paragraph." in result[0]["text"]
    assert "This is another paragraph." in result[0]["text"]
